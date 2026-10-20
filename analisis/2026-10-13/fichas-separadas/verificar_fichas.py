from __future__ import annotations

import importlib.util
import re
import sys
from pathlib import Path

from docx import Document
from pypdf import PdfReader


root = Path(r"F:\Tesis\App Tesis")
spec = importlib.util.spec_from_file_location("cards_model", root / "analisis/2026-10-13/fichas-separadas/generar_markdown.py")
assert spec and spec.loader
model = importlib.util.module_from_spec(spec)
sys.modules["cards_model"] = model
spec.loader.exec_module(model)

users = model.user_stories()
tech = model.tech_stories()
tasks = model.tasks()
assert (len(users), len(tech), len(tasks)) == (51, 19, 27)
assert len({item.code for item in users}) == 51
assert len({item.code for item in tech}) == 19
assert len({item.code for item in tasks}) == 27

markdown = {
    "hu": root / "documentacion/historias/usuario/historias_usuario_v2.5_2026-10-13.md",
    "ht": root / "documentacion/historias/tecnicas/historias_tecnicas_v1.5_2026-10-13.md",
    "pile": root / "documentacion/pila_producto/pila_producto_v1.5_2026-10-13.md",
    "sprint": root / "documentacion/sprints/sprint_backlog_v1.2_2026-10-13.md",
    "index": root / "documentacion/README.md",
    "change": root / "documentacion/control_de_cambios/registro_v2.7_2026-10-13.md",
    "plan": root / "documentacion/sprints/plan_sprints_v1.2_2026-10-13.md",
}
content = {name: path.read_text(encoding="utf-8") for name, path in markdown.items()}

assert len(re.findall(r"^### Historia \d+ —", content["hu"], re.M)) == 51
assert len(re.findall(r"^### Historia técnica \d+ —", content["ht"], re.M)) == 19
assert len(re.findall(r"^### Tarea \d+ —", content["sprint"], re.M)) == 27
base = content["pile"].split("## Línea base ordenada", 1)[1].split("## Diez historias técnicas en propuesta", 1)[0]
assert [int(n) for n in re.findall(r"^\| (\d+) \| `(?:HU|HT)-", base, re.M)] == list(range(1, 61))
proposal = content["pile"].split("## Diez historias técnicas en propuesta", 1)[1].split("## Pendientes para planificar", 1)[0]
assert len(re.findall(r"^\| \d+ \| `HT-", proposal, re.M)) == 10

for story in users:
    text = content["hu"]
    assert text.count(f"`{story.code}`") == 1, story.code
    assert story.need in text, story.code
    for criterion in story.criteria:
        assert criterion in text, (story.code, criterion)
for story in tech:
    assert content["ht"].count(f"| **Código de referencia** | `{story.code}` |") == 1, story.code
    assert story.exit.replace("\n", "<br>") in content["ht"], story.code
for task in tasks:
    assert content["sprint"].count(f"| **Código de referencia** | `{task.code}` |") == 1, task.code
    assert task.result in content["sprint"], task.code

for name, path in markdown.items():
    text = content[name]
    assert "\ufffd" not in text, path
    assert all(line == line.rstrip() for line in text.splitlines()), path
    for target in re.findall(r"\[[^]]+\]\(([^)]+)\)", text):
        if target.startswith(("http:", "https:")):
            continue
        assert (path.parent / target.split("#", 1)[0]).exists(), (path, target)

doc_stories = Document(root / "documentacion/historias/historias_usuario_tecnicas_pila_v1.2_2026-10-13.docx")
doc_sprints = Document(root / "documentacion/sprints/plan_incremental_y_sprint_backlog_v1.6_2026-10-13.docx")
stories_text = "\n".join(cell.text for table in doc_stories.tables for row in table.rows for cell in row.cells)
sprints_text = "\n".join(cell.text for table in doc_sprints.tables for row in table.rows for cell in row.cells)
assert all(story.code in stories_text for story in users + tech)
assert all(task.code in sprints_text for task in tasks)
assert len(PdfReader(root / "analisis/2026-10-13/fichas-separadas/historias_usuario_tecnicas_pila_v1.2_2026-10-13.pdf").pages) == 33
assert len(PdfReader(root / "analisis/2026-10-13/fichas-separadas/plan_incremental_y_sprint_backlog_v1.6_2026-10-13.pdf").pages) == 12
assert all(label in "\n".join(p.text for p in doc_sprints.paragraphs) for label in (
    "RESUMEN DE SPRINTS", "RELACIÓN CON HISTORIAS Y PILA", "EVIDENCIA Y CONDICIONES DE CIERRE"
))
for path in (root / "documentacion").rglob("*.docx"):
    assert not path.name.startswith("~$"), path
    assert Document(path).core_properties.author == "Franklin Acero", path

print("51 HU, 19 HT y 27 tareas conservadas en Markdown y Word")
print("60 posiciones base y 10 HT propuestas en la pila")
print("Criterios HU y salidas HT presentes; enlaces locales correctos")
print("Word renderizado: 33 páginas de historias y 12 de sprints")
