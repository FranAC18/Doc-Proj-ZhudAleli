from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


ROOT = Path(r"F:\Tesis\App Tesis")
HELPERS = ROOT / "analisis/2026-10-13/plan-incremental-docx/build_plan_v1_1.py"
SOURCE = ROOT / "documentacion/sprints/sprint_backlog_v1.0_2026-10-13.md"
OUTPUT = ROOT / "documentacion/sprints/plan_incremental_y_sprint_backlog_v1.2_2026-10-13.docx"

spec = importlib.util.spec_from_file_location("plan_helpers", HELPERS)
helpers = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(helpers)


def clean(value: str) -> str:
    value = re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value)
    return value.replace("**", "").replace("`", "").replace("<br>", "\n").strip()


def sprint_rows(number: int) -> list[list[str]]:
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    active = False
    rows: list[list[str]] = []
    for line in lines:
        if line.startswith(f"## S{number}  "):
            active = True
            continue
        if active and line.startswith("## "):
            break
        if active and line.startswith(f"| `SB-S{number}-"):
            rows.append([clean(x) for x in line.strip().strip("|").split("|")])
    assert rows, number
    return rows


def build() -> None:
    doc = Document(helpers.REFERENCE)
    for paragraph in list(doc.paragraphs)[15:]:
        paragraph._element.getparent().remove(paragraph._element)
    doc.paragraphs[6].text = "Plan incremental y sprint backlog"
    doc.paragraphs[8].text = "Aplicación web para la gestión organizacional del Comité de Mujeres Indígenas “Tránsito Amaguaña” de Zhud"
    doc.paragraphs[14].text = "13 de octubre de 2026"
    for index in (6, 8, 14):
        paragraph = doc.paragraphs[index]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.name = "Century Gothic"
            run.font.size = Pt(9 if index == 8 else 10)
            run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()
    title = doc.add_paragraph("Plan incremental y sprint backlog", style="Title")
    title.paragraph_format.space_after = Pt(12)
    helpers.body(doc, "Versión 1.2 · 13 de octubre de 2026 · Autor: Franklin Acero · Estado: revisión académica pendiente")
    helpers.heading(doc, "INTRODUCCIÓN")
    helpers.body(doc, "Este documento conserva los cinco módulos del plan incremental y los relaciona con las historias de usuario, las historias técnicas y la pila de producto. El Sprint 1 cubre requisitos y diseño del 15 de septiembre al 12 de octubre. El desarrollo empieza con el Sprint 2 el 13 de octubre y termina con el Sprint 6 el 21 de diciembre; el 22 y 23 se reservan para el cierre.")
    helpers.body(doc, "El sprint backlog es preliminar: las 51 historias de usuario y 9 técnicas existentes se distribuyen como candidatas, junto a 10 historias técnicas nuevas propuestas. No hay estimaciones del equipo ni capacidad verificada; ninguna fila se presenta como trabajo terminado o comprometido. Las decisiones abiertas y la revisión académica mantienen sus estados reales.")

    helpers.heading(doc, "CALENDARIO Y CORRESPONDENCIA")
    helpers.add_table(doc, ["Sprint", "Fechas de 2026", "Bloque del Gantt", "Hito previsto"], [
        ["S1", "15 sep–12 oct · 28 días", "Requisitos y diseño", "Línea base y puertas por verificar"],
        ["S2", "13–26 oct · 14 días", "Administración y acceso", "Primer recorrido privado probado"],
        ["S3", "27 oct–9 nov · 14 días", "Planificación", "Actividades y encargos probados"],
        ["S4", "10–23 nov · 14 días", "Asistencia y participación", "Asistencia y continuidad aprobada probadas"],
        ["S5", "24 nov–7 dic · 14 días", "Trazabilidad productiva", "Cadena productiva y comercio realmente incorporado"],
        ["S6", "8–21 dic · 14 días", "Reportes e integración", "Pruebas integradas y preparación de entrega"],
        ["Cierre", "22–23 dic · 2 días", "Validación y capacitación", "Estado final con evidencia"],
    ], [1.7, 3.4, 4.2, 5.7])
    helpers.body(doc, "El plan completo del 15 de septiembre al 23 de diciembre suma 100 días calendario, 12 menos que las 16 semanas del Gantt original. La etiqueta S1 de SPEC-001 y de su plan técnico aprobado corresponde al bloque de acceso que este calendario llama S2; esos artefactos históricos no se renumeran sin revisión.")

    helpers.heading(doc, "PRIORIDAD Y REGLA DE ENTRADA")
    helpers.add_table(doc, ["Nivel", "Uso"], [
        ["Crítica", "Bloquea un recorrido esencial, permisos, integridad o evidencia de calidad."],
        ["Alta", "Condición importante para el sprint o reducción de un riesgo importante."],
        ["Media", "Necesaria para una entrega posterior, pero sin bloquear el mínimo actual."],
        ["Baja", "Mejora diferible; ninguna HT actual recibió este nivel."],
    ], [2.5, 12.5])
    helpers.body(doc, "La posición 1/60 de la pila expresa secuencia propuesta, no gravedad. Solo se selecciona una historia cuando hay requisito y actor, regla resuelta, dependencia disponible, criterios comprobables, tamaño estimado y responsable con tiempo acordado. Las 10 HT adicionales se mantienen como propuestas separadas de la línea base de 60.")

    sprint_info = {
        1: ("REQUISITOS Y DISEÑO", "Consolidar fuentes, validar reglas, preparar arquitectura, recorridos y pila. Los documentos existentes son evidencia parcial; la revisión y varias decisiones siguen pendientes."),
        2: ("ADMINISTRACIÓN Y ACCESO", "Demostrar el primer recorrido privado con permisos comprobados en servidor. SPEC-001 conserva la etiqueta histórica S1."),
        3: ("PLANIFICACIÓN", "Registrar actividades, responsables y encargos; verificar que el acceso anterior sigue funcionando."),
        4: ("ASISTENCIA Y CONTINUIDAD", "Registrar asistencia y corregir con historial. El QR y las operaciones sin red dependen de P-02, P-13 y P-14; P-03 limita sanciones."),
        5: ("TRAZABILIDAD Y RENDICIÓN", "Probar primero la cadena materia prima–secado–lote–existencias. Feria, ventas y saldos entran según reglas y capacidad."),
        6: ("REPORTES E INTEGRACIÓN", "Producir reportes de datos existentes, probar recorridos integrados y preparar validación, manuales y capacitación."),
    }
    for number in range(1, 7):
        label, purpose = sprint_info[number]
        helpers.heading(doc, f"SPRINT {number}  {label}")
        helpers.body(doc, purpose, keep_next=True)
        headers = ["Tarea", "Historia de usuario", "Historia técnica", "Prueba o entregable"]
        if number == 1:
            headers = ["Tarea", "Historia técnica", "Entregable", "Estado y prioridad"]
        helpers.add_table(doc, headers, sprint_rows(number), [3.0, 3.0, 4.0, 5.0], size=7.5)
        if number == 3:
            helpers.body(doc, "HU-CAT-001 conserva su bloque principal en S5. Solo se prepara el catálogo mínimo necesario para HU-PLA-005 si se aprueba y estima.")
        if number == 5:
            helpers.body(doc, "S5 concentra 23 HU candidatas. La cadena mínima se prioriza; el bloque comercial requiere estimación y decisiones. La fecha no obliga a declarar completo todo el grupo.")

    helpers.heading(doc, "PRUEBAS Y CIERRE")
    helpers.body(doc, "Las pruebas parciales acompañan S2–S6. Cada incremento conserva permisos, integridad, uso móvil, pruebas de regresión y documentación. La integración comienza al conectar módulos y se revisa de forma final en S6. La validación con usuarias, capacitación y manuales se preparan durante S6; el cierre del 22 y 23 de diciembre registra lo realizado y lo pendiente.")
    helpers.body(doc, "En la revisión de cada sprint se anotan historias completas, parciales y bloqueadas, pruebas con resultados reales, observaciones, esfuerzo restante y una retrospectiva. Cambios de fecha, recursos o alcance requieren autorización y, cuando afectan el anteproyecto, revisión de la directora.")

    helpers.heading(doc, "FUENTES E HISTORIAL")
    helpers.body(doc, "Fuentes: Gantt aportado por el tesista; catálogo CAT-REQ-001 v1.1; historias de usuario v2.3; historias técnicas v1.4; pila de producto v1.3; plan de sprints v1.1; sprint backlog v1.0; constitución 0.2.0 y ADR-0006. El formato visual procede de 7402_alcance.docx.")
    helpers.body(doc, "Historial: v1.0, plan inicial de cinco incrementos; v1.1, adaptación al Gantt de 16 semanas; v1.2, S1 de requisitos y diseño, desarrollo S2–S6 y tareas relacionadas con HU, HT y pila. Publicación pendiente de confirmación.")
    ending = doc.add_paragraph()
    ending.paragraph_format.space_after = Pt(0)
    ending.paragraph_format.line_spacing = Pt(1)
    ending.add_run(" ").font.size = Pt(1)

    doc.core_properties.title = "Plan incremental y sprint backlog"
    doc.core_properties.subject = "Zhud Alelí"
    doc.core_properties.author = "Franklin Acero"
    doc.save(OUTPUT)
    print(OUTPUT)
    print("filas", {n: len(sprint_rows(n)) for n in range(1, 7)})


if __name__ == "__main__":
    build()
