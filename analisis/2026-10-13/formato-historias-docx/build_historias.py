from __future__ import annotations

import html
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor


ROOT = Path(r"F:\Tesis\App Tesis")
REFERENCE = Path(r"F:\Software\8vo Semestre\Aplicaciones Informaticas II\7402_alcance.docx")
OUTPUT = ROOT / "documentacion" / "historias" / "historias_usuario_tecnicas_pila_v1.0_2026-10-13.docx"
HU_SOURCE = ROOT / "documentacion" / "historias" / "usuario" / "historias_usuario_v2.2_2026-10-12.md"
HT_SOURCE = ROOT / "documentacion" / "historias" / "tecnicas" / "historias_tecnicas_v1.3_2026-10-12.md"
PB_SOURCE = ROOT / "documentacion" / "pila_producto" / "pila_producto_v1.2_2026-10-12.md"


def cells(line: str) -> list[str]:
    return [x.strip() for x in line.strip().strip("|").split("|")]


def clean(value: str) -> str:
    value = html.unescape(value)
    value = re.sub(r"<br\s*/?>", "\n", value, flags=re.I)
    value = value.replace("**", "").replace("`", "")
    value = re.sub(r"<[^>]+>", "", value)
    return value.strip()


def parse_rows(path: Path, start: str, stop: str | None = None) -> list[list[str]]:
    lines = path.read_text(encoding="utf-8").splitlines()
    active = False
    rows: list[list[str]] = []
    for line in lines:
        if line.startswith(start):
            active = True
            continue
        if active and stop and line.startswith(stop):
            break
        if active and line.startswith("|") and not line.startswith("|---") and not line.startswith("| :"):
            row = cells(line)
            if row[0] in {"Historia y requisito", "Historia técnica", "Orden", "Historia técnica "}:
                continue
            if row[0].startswith("---"):
                continue
            rows.append([clean(c) for c in row])
    return rows


def get_user_stories() -> list[tuple[str, list[list[str]]]]:
    lines = HU_SOURCE.read_text(encoding="utf-8").splitlines()
    groups: list[tuple[str, list[list[str]]]] = []
    current: list[list[str]] | None = None
    for line in lines:
        if line.startswith("## "):
            heading = line[3:].strip()
            if heading.startswith("Sprint ") or heading == "Continuidad transversal":
                current = []
                groups.append((heading, current))
            elif heading == "Condiciones de la edición":
                break
        elif current is not None and line.startswith("|"):
            row = cells(line)
            if len(row) == 4 and "HU-" in row[0]:
                current.append([clean(c) for c in row])
    assert sum(len(x[1]) for x in groups) == 51, [(h, len(r)) for h, r in groups]
    return groups


def get_technical_stories() -> list[tuple[list[str], list[str]]]:
    summary = parse_rows(HT_SOURCE, "## Resumen de trabajo", "## Criterios, tareas y pruebas")
    details = parse_rows(HT_SOURCE, "## Criterios, tareas y pruebas", "## Límites de esta propuesta")
    summary = [r for r in summary if r and r[0].startswith("HT-")]
    details = [r for r in details if r and r[0].startswith("HT-")]
    dmap = {re.search(r"HT-GEN-\d{3}", r[0]).group(): r for r in details}
    assert len(summary) == len(details) == 9, (len(summary), len(details))
    return [(s, dmap[re.search(r"HT-GEN-\d{3}", s[0]).group()]) for s in summary]


def get_backlog() -> list[list[str]]:
    lines = PB_SOURCE.read_text(encoding="utf-8").splitlines()
    rows: list[list[str]] = []
    for line in lines:
        if line.startswith("| "):
            row = cells(line)
            if len(row) == 6 and row[0].isdigit():
                rows.append([clean(c) for c in row])
    assert len(rows) == 60, len(rows)
    assert [int(r[0]) for r in rows] == list(range(1, 61))
    return rows


def set_cell_shading(cell, fill: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcPr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_border(cell) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    borders = tcPr.find(qn("w:tcBorders"))
    if borders is None:
        borders = OxmlElement("w:tcBorders")
        tcPr.append(borders)
    for edge in ("top", "left", "bottom", "right"):
        el = borders.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            borders.append(el)
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:color"), "D9D9D9")


def set_cell_padding(cell, top: int = 70, bottom: int = 70, left: int = 95, right: int = 95) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    mar = tcPr.find(qn("w:tcMar"))
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcPr.append(mar)
    for edge, amount in (("top", top), ("bottom", bottom), ("left", left), ("right", right)):
        el = mar.find(qn(f"w:{edge}"))
        if el is None:
            el = OxmlElement(f"w:{edge}")
            mar.append(el)
        el.set(qn("w:w"), str(amount))
        el.set(qn("w:type"), "dxa")


def make_header_row(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    trPr.append(header)


def no_split(row) -> None:
    trPr = row._tr.get_or_add_trPr()
    trPr.append(OxmlElement("w:cantSplit"))


def put(cell, value: str, *, bold: bool = False, white: bool = False, center: bool = False, size: float = 8) -> None:
    cell.text = ""
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    set_cell_border(cell)
    set_cell_padding(cell)
    for idx, line in enumerate(value.split("\n")):
        p = cell.paragraphs[0] if idx == 0 else cell.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER if center else WD_ALIGN_PARAGRAPH.LEFT
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2 if idx < len(value.split("\n")) - 1 else 0)
        p.paragraph_format.line_spacing = 1.1
        run = p.add_run(line)
        run.bold = bold
        run.font.name = "Century Gothic"
        run.font.size = Pt(size)
        run.font.color.rgb = RGBColor(255, 255, 255) if white else RGBColor(0, 0, 0)


def story_table(doc, title: str, rows: list[tuple[str, str]]) -> None:
    table = doc.add_table(rows=1, cols=2)
    table.autofit = False
    table.columns[0].width = Cm(3.1)
    table.columns[1].width = Cm(11.9)
    header = table.rows[0]
    merged = header.cells[0].merge(header.cells[1])
    put(merged, title, bold=True, white=True, size=8.5)
    set_cell_shading(merged, "404040")
    make_header_row(header)
    no_split(header)
    for index, (label, value) in enumerate(rows):
        row = table.add_row()
        row.cells[0].width = Cm(3.1)
        row.cells[1].width = Cm(11.9)
        put(row.cells[0], label, bold=True)
        put(row.cells[1], value)
        if index % 2:
            for c in row.cells:
                set_cell_shading(c, "F7F7F7")
        no_split(row)
    # A single story is short enough to stay intact on an A4 page.
    for row in table.rows[:-1]:
        for cell in row.cells:
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.keep_with_next = True
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)


def add_body(doc, text: str):
    p = doc.add_paragraph(style="Normal")
    p.paragraph_format.space_after = Pt(7)
    p.add_run(text)
    return p


def add_heading(doc, text: str, level: int = 1) -> None:
    p = doc.add_paragraph(text, style=f"Heading {level}")
    p.paragraph_format.keep_with_next = True


def build() -> None:
    groups = get_user_stories()
    technical = get_technical_stories()
    backlog = get_backlog()

    doc = Document(REFERENCE)
    # Retain the cover and the source's institutional images, header and footer.
    for p in list(doc.paragraphs)[15:]:
        p._element.getparent().remove(p._element)
    doc.paragraphs[6].text = "Historias de usuario, historias técnicas y pila de producto"
    doc.paragraphs[8].text = "Aplicación web para la gestión organizacional del Comité de Mujeres Indígenas “Tránsito Amaguaña” de Zhud"
    doc.paragraphs[14].text = "13 de octubre de 2026"
    for idx in (6, 8, 14):
        p = doc.paragraphs[idx]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.bold = True
            run.font.name = "Century Gothic"
            run.font.size = Pt(10 if idx != 8 else 9)
            run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()
    title = doc.add_paragraph("Historias de usuario, historias técnicas y pila de producto", style="Title")
    title.paragraph_format.space_after = Pt(12)
    add_body(doc, "Versión 1.0 · 13 de octubre de 2026 · Estado: edición documental para revisión académica")

    add_heading(doc, "INTRODUCCIÓN")
    add_body(doc, "Este documento reúne las 51 historias de usuario, las 9 historias técnicas y los 60 elementos de la pila de producto de Zhud Alelí. Sirve para relacionar las necesidades de la organización con los criterios de aceptación y el orden preliminar de trabajo.")
    add_body(doc, "Las historias se basan en el catálogo CAT-REQ-001 v1.1 y en los documentos formales citados al final. El orden de prioridad es una propuesta; las estimaciones siguen pendientes del equipo. Las ideas surgidas de la entrevista que todavía requieren validación no se mezclan con esta línea base. Ninguna historia se presenta como implementada o probada.")

    add_heading(doc, "HISTORIAS DE USUARIO")
    add_body(doc, "Cada historia expresa quién necesita una función, qué espera lograr y cómo se comprobará su cumplimiento. La agrupación por sprint indica la etapa prevista y puede ajustarse durante el refinamiento.")
    for heading, stories in groups:
        add_heading(doc, heading, 2)
        for row in stories:
            identity, need, acceptance, planning = row
            title_bits = identity.split("\n", 1)
            story_table(doc, title_bits[0], [
                ("Requisito", title_bits[1] if len(title_bits) > 1 else ""),
                ("Necesidad", need),
                ("Aceptación", acceptance),
                ("Planificación", planning),
            ])

    add_heading(doc, "HISTORIAS TÉCNICAS")
    technical_intro = add_body(doc, "Estas historias preparan las capacidades técnicas necesarias para cumplir las historias de usuario. Sus pruebas se realizarán cuando exista un incremento funcional; sus estimaciones permanecen pendientes.")
    technical_intro.paragraph_format.keep_with_next = True
    for summary, detail in technical:
        identity, kind_description, priority, deps = summary
        _, criteria, tasks, test = detail
        story_table(doc, identity, [
            ("Tipo y propósito", kind_description),
            ("Criterios de aceptación", criteria),
            ("Tareas previstas", tasks),
            ("Prueba prevista", test),
            ("Planificación", priority + "\nDependencias: " + deps),
        ])

    add_heading(doc, "PILA DE PRODUCTO")
    add_body(doc, "La posición 1 representa la prioridad propuesta más alta. La secuencia combina dependencias, valor de los recorridos y etapas previstas; no supone que todos los elementos de una etapa caben en un sprint. Los 60 elementos se reordenarán tras el refinamiento y la validación que corresponda.")
    table = doc.add_table(rows=1, cols=4)
    table.autofit = False
    for col, width in zip(table.columns, (Cm(1.0), Cm(5.6), Cm(5.0), Cm(3.4))):
        col.width = width
    for c, label in zip(table.rows[0].cells, ("Orden", "Historia y elemento", "Etapa y dependencias", "Preparación")):
        put(c, label, bold=True, white=True, center=True, size=8)
        set_cell_shading(c, "404040")
    make_header_row(table.rows[0])
    no_split(table.rows[0])
    for i, row in enumerate(backlog):
        order, id_type, element, stage, deps, prep = row
        tr = table.add_row()
        for c, width in zip(tr.cells, (Cm(1.0), Cm(5.6), Cm(5.0), Cm(3.4))):
            c.width = width
        contents = (order, id_type + "\n" + element, stage + "\n" + deps, prep)
        for j, (c, value) in enumerate(zip(tr.cells, contents)):
            put(c, value, center=(j == 0), size=7.5 if j else 8)
            if i % 2:
                set_cell_shading(c, "F7F7F7")
        no_split(tr)

    add_heading(doc, "CONDICIONES DE PLANIFICACIÓN")
    add_body(doc, "Antes de incorporar una historia a un sprint deben revisarse sus dependencias, resolver los puntos abiertos que afecten su comportamiento y asignar una estimación realista. Las decisiones pendientes P-02 a P-14 se mantienen como límites del alcance afectado. Los requisitos y reglas derivados de la entrevista requieren su propio proceso de validación antes de incorporarse a esta pila.")
    add_body(doc, "La revisión académica de este documento y la validación organizacional de los ajustes posteriores a CAT-REQ-001 v1.1 figuran pendientes en las fuentes. Esta edición no cambia ese estado.")

    add_heading(doc, "FUENTES Y CONTROL DE CAMBIOS")
    sources = [
        "Catálogo de requisitos CAT-REQ-001 v1.1.",
        "Historias de usuario v2.2, 12 de octubre de 2026.",
        "Historias técnicas v1.3, 12 de octubre de 2026.",
        "Pila de producto v1.2, 12 de octubre de 2026.",
        "Formato visual: 7402_alcance.docx, entregado como referencia por el tesista.",
    ]
    for item in sources:
        add_body(doc, item)
    add_body(doc, "Versión 1.0, 13 de octubre de 2026: se integran 51 historias de usuario, 9 historias técnicas y la pila de 60 elementos en un documento con el formato académico de referencia.")

    doc.core_properties.title = "Historias de usuario, historias técnicas y pila de producto"
    doc.core_properties.subject = "Zhud Alelí"
    doc.core_properties.author = "Franklin Acero"
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)
    print("HU", sum(len(s) for _, s in groups), "HT", len(technical), "PB", len(backlog))


if __name__ == "__main__":
    build()
