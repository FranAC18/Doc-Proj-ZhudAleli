from __future__ import annotations

import importlib.util
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor


ROOT = Path(r"F:\Tesis\App Tesis")
HELPERS = ROOT / "analisis/2026-10-13/plan-incremental-docx/build_plan_v1_1.py"
SPRINT_SOURCE = ROOT / "documentacion/sprints/sprint_backlog_v1.1_2026-10-13.md"
HU_SOURCE = ROOT / "documentacion/historias/usuario/historias_usuario_v2.4_2026-10-13.md"
OUTPUT = ROOT / "documentacion/sprints/plan_incremental_y_sprint_backlog_v1.3_2026-10-13.docx"

spec = importlib.util.spec_from_file_location("plan_helpers", HELPERS)
helpers = importlib.util.module_from_spec(spec)
assert spec and spec.loader
spec.loader.exec_module(helpers)


def clean(value: str) -> str:
    return re.sub(r"\[([^]]+)\]\([^)]+\)", r"\1", value).replace("**", "").replace("`", "").replace("<br>", "\n").strip()


def sprint_rows(number: int) -> list[list[str]]:
    text = SPRINT_SOURCE.read_text(encoding="utf-8")
    section = text.split(f"## S{number}  ", 1)[1].split("## ", 1)[0]
    rows = []
    for line in section.splitlines():
        if not re.match(rf"^\| \*\*(?:Esencial|Alta|Complementaria)\*\* \| `SB-S{number}-", line):
            continue
        parts = [clean(part) for part in line.strip().strip("|").split("|")]
        assert len(parts) == 3
        task_lines = parts[1].splitlines()
        task_id, title = task_lines[0].split(" ", 1)
        links = "\n".join(task_lines[1:])
        rows.append([f"{task_id}\n{parts[0]}", f"{title}\n{links}", parts[2]])
    assert rows, number
    return rows


def hu_rows() -> dict[str, list[list[str]]]:
    text = HU_SOURCE.read_text(encoding="utf-8")
    sections = {}
    mapping = [
        ("S2", "## Sprint 2"), ("S3", "## Sprint 3"),
        ("S4", "## Sprint 4"), ("S5", "## Sprint 5"),
        ("S6", "## Sprint 6"), ("Transversal", "## Continuidad transversal"),
    ]
    for label, heading in mapping:
        section = text.split(heading, 1)[1].split("### Criterios y dependencias", 1)[0]
        rows = []
        for line in section.splitlines():
            if not re.match(r"^\| \*\*(?:Esencial|Alta|Complementaria)\*\* \| `HU-", line):
                continue
            parts = [clean(part) for part in line.strip().strip("|").split("|")]
            assert len(parts) == 4
            hu_id, title = parts[1].split(" · ", 1)
            rows.append([hu_id, title, parts[0], parts[3]])
        sections[label] = rows
    assert sum(map(len, sections.values())) == 51
    return sections


def build() -> None:
    doc = Document(helpers.REFERENCE)
    for paragraph in list(doc.paragraphs)[15:]:
        paragraph._element.getparent().remove(paragraph._element)
    doc.paragraphs[6].text = "Plan incremental y sprint backlog"
    doc.paragraphs[8].text = "Aplicación web para la gestión organizacional del Comité de Mujeres Indígenas “Tránsito Amaguaña” de Zhud"
    doc.paragraphs[14].text = "13 de octubre de 2026"
    for index in (6, 8, 14):
        paragraph = doc.paragraphs[index]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in paragraph.runs:
            run.bold = True
            run.font.name = "Century Gothic"
            run.font.size = Pt(9 if index == 8 else 10)
            run.font.color.rgb = RGBColor(0, 0, 0)

    doc.add_page_break()
    title = doc.add_paragraph("Plan incremental y sprint backlog", style="Title")
    title.paragraph_format.space_after = Pt(10)
    helpers.body(doc, "Versión 1.3 · 13 de octubre de 2026 · Autor: Franklin Acero · Estado: revisión académica pendiente")
    helpers.heading(doc, "INTRODUCCIÓN")
    helpers.body(doc, "Este documento organiza el calendario, las tareas candidatas de cada sprint y la importancia propuesta de las historias de usuario. S1 corresponde a requisitos y diseño del 15 de septiembre al 12 de octubre. El desarrollo empieza con S2 el 13 de octubre y el cierre se reserva para el 22 y 23 de diciembre. Las tablas separan prioridad, trabajo, vínculos y evidencia de cierre.")
    helpers.body(doc, "La posición 1/60 en la pila señala secuencia propuesta, no importancia. La importancia de las HU se hereda de sus requisitos: Esencial, Alta o Complementaria. La prioridad de cada tarea del sprint también es una propuesta; depende del valor y de las puertas técnicas. No hay estimaciones ni capacidad verificadas y ninguna fila se presenta como trabajo terminado.")

    helpers.heading(doc, "CALENDARIO GENERAL")
    helpers.add_table(doc, ["Etapa", "Fechas de 2026", "Objetivo principal"], [
        ["S1", "15 sep–12 oct", "Requisitos y diseño revisables"],
        ["S2", "13–26 oct", "Administración y acceso"],
        ["S3", "27 oct–9 nov", "Planificación y encargos"],
        ["S4", "10–23 nov", "Asistencia y continuidad aprobada"],
        ["S5", "24 nov–7 dic", "Trazabilidad productiva y rendición"],
        ["S6", "8–21 dic", "Reportes, integración y preparación de cierre"],
        ["Cierre", "22–23 dic", "Validación, capacitación y entrega según resultado real"],
    ], [2.0, 3.8, 9.2], size=8.2)
    helpers.body(doc, "El periodo completo suma 100 días calendario. La etiqueta S1 del plan técnico de SPEC-001 es histórica y corresponde al bloque de acceso que este calendario denomina S2. Las decisiones pendientes y la capacidad real pueden modificar la selección de trabajo, con registro y autorización.")

    names = {
        1: "REQUISITOS Y DISEÑO", 2: "ADMINISTRACIÓN Y ACCESO",
        3: "PLANIFICACIÓN", 4: "ASISTENCIA Y CONTINUIDAD",
        5: "TRAZABILIDAD Y RENDICIÓN", 6: "REPORTES E INTEGRACIÓN",
    }
    purposes = {
        1: "Consolidar fuentes y hacer visible qué falta validar antes del desarrollo. No se acreditan aprobaciones por la fecha.",
        2: "Preparar el primer recorrido privado y comprobar permisos en el servidor.",
        3: "Registrar actividades, responsables y encargos y mantener las pruebas de acceso.",
        4: "Registrar asistencia con historial y probar operaciones sin red solo si sus decisiones están cerradas.",
        5: "Probar primero la cadena de materia prima, lotes y existencias; incluir comercio según reglas y capacidad.",
        6: "Generar reportes de datos existentes, ejecutar integración y preparar validación y manuales.",
    }
    for number in range(1, 7):
        helpers.heading(doc, f"SPRINT {number}  {names[number]}")
        helpers.body(doc, purposes[number], keep_next=True)
        helpers.add_table(doc, ["Tarea y prioridad", "Trabajo e historias relacionadas", "Cierre verificable"], sprint_rows(number), [2.5, 6.7, 5.8], size=8.0)
        if number == 3:
            helpers.body(doc, "HU-CAT-001 conserva su bloque principal en S5. Se adelanta únicamente la parte mínima necesaria para HU-PLA-005, si se aprueba y estima.")
        if number == 4:
            helpers.body(doc, "P-03 mantiene pendiente el procedimiento de justificativos y sanciones; el sprint no presume multas automáticas. P-02, P-13 y P-14 condicionan QR y operación sin red.")
        if number == 5:
            helpers.body(doc, "S5 reúne 23 HU candidatas. Las tareas de feria, venta y rendición se seleccionan después de estimar y resolver sus reglas; la fecha no demuestra que todo el grupo esté terminado.")

    helpers.heading(doc, "HISTORIAS DE USUARIO Y SU IMPORTANCIA")
    helpers.body(doc, "Las siguientes tablas permiten encontrar cada HU por sprint, nivel de importancia y posición en la pila. Las 51 historias mantienen sus IDs, requisitos y criterios completos en el documento de historias de usuario v2.4. El orden visible aquí agrupa Esencial, Alta y Complementaria; no sustituye la secuencia de dependencias.")
    for label, rows in hu_rows().items():
        helpers.heading(doc, f"HISTORIAS {label.upper()}", 2)
        helpers.add_table(doc, ["ID", "Historia", "Importancia", "Pila"], rows, [2.6, 7.4, 3.1, 1.9], size=8.0)

    helpers.heading(doc, "PRUEBAS Y CIERRE")
    helpers.body(doc, "Las pruebas parciales acompañan S2–S6. Cada revisión registra historias completas, parciales y bloqueadas, resultados reales, defectos, regresión y esfuerzo restante. La integración se prueba conforme se conectan recorridos. El cierre del 22 y 23 de diciembre registra la validación con usuarias y la capacitación realizadas o pendientes, sin presentarlas como ejecutadas antes de tiempo.")
    helpers.heading(doc, "FUENTES E HISTORIAL")
    helpers.body(doc, "Fuentes: Gantt y fechas aportados por el tesista; CAT-REQ-001 v1.1; priorización de requisitos v1.2; historias de usuario v2.4; historias técnicas v1.4; pila de producto v1.4; plan de sprints v1.2; sprint backlog v1.1; constitución 0.2.0 y ADR-0006. Formato visual basado en 7402_alcance.docx.")
    helpers.body(doc, "Historial: v1.2, calendario S1–S6 y backlog preliminar. v1.3, tablas de sprint con prioridad visible y anexo de 51 HU con importancia y posición en pila. Revisión académica y publicación pendientes.")
    ending = doc.add_paragraph()
    ending.paragraph_format.space_after = Pt(0)
    ending.paragraph_format.line_spacing = Pt(1)
    ending.add_run(" ").font.size = Pt(1)
    doc.core_properties.title = "Plan incremental y sprint backlog"
    doc.core_properties.subject = "Zhud Alelí"
    doc.core_properties.author = "Franklin Acero"
    doc.save(OUTPUT)
    print(OUTPUT)
    print("tareas", {n: len(sprint_rows(n)) for n in range(1, 7)})
    print("HU", {name: len(rows) for name, rows in hu_rows().items()})


if __name__ == "__main__":
    build()
