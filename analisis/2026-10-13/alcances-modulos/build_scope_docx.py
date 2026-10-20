from __future__ import annotations
import re
from pathlib import Path
from datetime import date
from html import unescape
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

ROOT = Path(r"F:\Tesis\App Tesis")
SOURCE = ROOT / "documentacion" / "alcance" / "alcance_v1.3_2026-10-13.md"
TEMPLATE = ROOT / "analisis" / "2026-10-13" / "alcances-modulos" / "Formato Corregido.docx"
OUTPUT = ROOT / "documentacion" / "alcance" / "alcance_v1.3_2026-10-13.docx"
GREEN = "4C5B4A"
BURGUNDY = "642138"
TERRACOTTA = "AC6E40"
CREAM = "F0EAC8"
INK = "292929"
CONTENT_WIDTH = 6.87


def parse_frontmatter(text: str):
    match = re.match(r"\A---\s*\n(.*?)\n---\s*\n", text, re.S)
    if not match:
        raise ValueError("No se encontró el frontmatter del Markdown de origen")
    metadata = {}
    for line in match.group(1).splitlines():
        if ":" not in line:
            continue
        key, value = line.split(":", 1)
        metadata[key.strip()] = value.strip().strip('"')
    return metadata, text[match.end():]


def parse_table(lines, start):
    table_lines = []
    i = start
    while i < len(lines) and lines[i].strip().startswith("|"):
        table_lines.append(lines[i].strip())
        i += 1
    rows = []
    for raw in table_lines:
        cells = [c.strip() for c in raw.strip("|").split("|")]
        if cells and all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) for c in cells):
            continue
        rows.append(cells)
    return rows, i


def strip_code_fence(text):
    return text.replace(chr(96), "")


def add_hyperlink(paragraph, label, target, bold=False, italic=False):
    rid = paragraph.part.relate_to(target, RT.HYPERLINK, is_external=True)
    hyperlink = OxmlElement("w:hyperlink")
    hyperlink.set(qn("r:id"), rid)
    run = OxmlElement("w:r")
    rpr = OxmlElement("w:rPr")
    color = OxmlElement("w:color")
    color.set(qn("w:val"), BURGUNDY)
    rpr.append(color)
    fonts = OxmlElement("w:rFonts")
    fonts.set(qn("w:ascii"), "Heebo Light")
    fonts.set(qn("w:hAnsi"), "Heebo Light")
    rpr.append(fonts)
    if bold:
        rpr.append(OxmlElement("w:b"))
    if italic:
        rpr.append(OxmlElement("w:i"))
    run.append(rpr)
    text_node = OxmlElement("w:t")
    text_node.text = label
    run.append(text_node)
    hyperlink.append(run)
    paragraph._p.append(hyperlink)


INLINE_RE = re.compile(r"\[([^\]]+)\]\(([^)]+)\)|\*\*(.+?)\*\*|`([^`]+)`|(?<!\*)\*([^*]+)\*(?!\*)")


def set_run(run, *, bold=False, italic=False, code=False, size=None, color=INK):
    run.bold = bold
    run.italic = italic
    run.font.name = "Consolas" if code else "Heebo Light"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), run.font.name)
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), run.font.name)
    run.font.color.rgb = RGBColor.from_string(color)
    if size is not None:
        run.font.size = Pt(size)


def add_inline(paragraph, text, size=None, color=INK):
    text = unescape(text)
    pos = 0
    for match in INLINE_RE.finditer(text):
        if match.start() > pos:
            run = paragraph.add_run(text[pos:match.start()])
            set_run(run, size=size, color=color)
        label, target, bold_text, code_text, italic_text = match.groups()
        if label is not None:
            add_hyperlink(paragraph, label, target)
        else:
            value = bold_text if bold_text is not None else (code_text if code_text is not None else italic_text)
            run = paragraph.add_run(value)
            set_run(run, bold=bold_text is not None, italic=italic_text is not None, code=code_text is not None, size=size, color=color)
        pos = match.end()
    if pos < len(text):
        run = paragraph.add_run(text[pos:])
        set_run(run, size=size, color=color)


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)
    shd.set(qn("w:val"), "clear")


def set_cell_margins(cell, top=70, start=85, bottom=70, end=85):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    margins = tc_pr.first_child_found_in("w:tcMar")
    if margins is None:
        margins = OxmlElement("w:tcMar")
        tc_pr.append(margins)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{side}"))
        if node is None:
            node = OxmlElement(f"w:{side}")
            margins.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_table_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = tbl_pr.find(qn("w:tblBorders"))
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = qn(f"w:{edge}")
        element = borders.find(tag)
        if element is None:
            element = OxmlElement(f"w:{edge}")
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), "6")
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), TERRACOTTA)


def set_repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def choose_widths(headers, rows, kind="body"):
    n = len(headers)
    if kind == "metadata":
        return [2.12, CONTENT_WIDTH - 2.12]
    if kind == "revisions":
        return [0.65, 0.9, 3.4, CONTENT_WIDTH - 4.95]
    normalized = [re.sub(r"[^a-z]", "", h.lower()) for h in headers]
    if normalized[:3] == ["bloquedealcance", "funcionalidadincluida", "requisitosfuncionales"]:
        return [1.38, 2.86, 1.55, CONTENT_WIDTH - 5.79]
    if normalized == ["decisin", "temaabierto", "lmitequedebeconservarse"]:
        return [0.95, 2.15, CONTENT_WIDTH - 3.10]
    if n == 1:
        return [CONTENT_WIDTH]
    if n == 2:
        return [1.45, CONTENT_WIDTH - 1.45]
    if n == 3:
        return [1.5, 2.15, CONTENT_WIDTH - 3.65]
    if n == 4:
        return [1.25, 2.55, 1.25, CONTENT_WIDTH - 5.05]
    weights = [1.0] * n
    return [CONTENT_WIDTH * w / sum(weights) for w in weights]


def add_table(doc, rows, kind="body"):
    if not rows:
        return None
    ncols = max(len(row) for row in rows)
    normalized = [list(row) + [""] * (ncols - len(row)) for row in rows]
    table = doc.add_table(rows=len(normalized), cols=ncols)
    table.autofit = False
    table.alignment = 0
    widths = choose_widths(normalized[0], normalized[1:], kind)
    tbl_pr = table._tbl.tblPr
    layout = tbl_pr.find(qn("w:tblLayout"))
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")
    for col_idx, width in enumerate(widths):
        table.columns[col_idx].width = Inches(width)
        for row in table.rows:
            row.cells[col_idx].width = Inches(width)
    set_table_borders(table)
    if kind == "body":
        set_repeat_table_header(table.rows[0])
    for ri, row_data in enumerate(normalized):
        row = table.rows[ri]
        for ci, raw in enumerate(row_data):
            cell = row.cells[ci]
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
            set_cell_margins(cell)
            if kind == "metadata" and ri > 0 and ci == 0:
                set_cell_shading(cell, BURGUNDY)
                font_color = CREAM
                bold = True
            elif ri == 0:
                set_cell_shading(cell, GREEN)
                font_color = CREAM
                bold = True
            else:
                set_cell_shading(cell, CREAM if ri % 2 == 0 else "FFFFFF")
                font_color = INK
                bold = False
            p = cell.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(1)
            p.paragraph_format.line_spacing = 1.02
            if ri == 0:
                p.paragraph_format.keep_with_next = True
            add_inline(p, str(raw), size=8.25 if ncols >= 3 else 8.8, color=font_color)
            for run in p.runs:
                run.bold = bold or run.bold
                if bold:
                    run.font.color.rgb = RGBColor.from_string(font_color)
        cant_split = OxmlElement("w:cantSplit")
        row._tr.get_or_add_trPr().append(cant_split)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table


def add_heading(doc, text, level):
    style = "Heading 1" if level == 1 else ("Heading 2" if level == 2 else "Heading 3")
    p = doc.add_paragraph(style=style)
    p.paragraph_format.keep_with_next = True
    p.paragraph_format.space_before = Pt(10 if level == 1 else 6)
    p.paragraph_format.space_after = Pt(3.5)
    add_inline(p, text, size={1: 15, 2: 12, 3: 10.5}[level], color=GREEN if level == 1 else BURGUNDY)
    for run in p.runs:
        run.bold = True
        run.font.name = "Georgia"
        run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Georgia")
        run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Georgia")
    return p


def parse_source_body(text):
    lines = text.splitlines()
    start = next((i for i, line in enumerate(lines) if re.match(r"^## 1\. Propósito", line)), None)
    if start is None:
        raise ValueError("No se encontró el comienzo del contenido del alcance")
    lines = lines[start:]
    return lines


def convert_body(doc, lines):
    i = 0
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1
            continue
        if line.startswith("|"):
            rows, i = parse_table(lines, i)
            add_table(doc, rows)
            continue
        heading = re.match(r"^(#{1,6})\s+(.*)$", line)
        if heading:
            hashes, title = heading.groups()
            level = min(len(hashes) - 1, 3)
            if level >= 1:
                heading_paragraph = add_heading(doc, title, level)
                if title.strip() == "Historial de cambios":
                    heading_paragraph.paragraph_format.page_break_before = True
            i += 1
            continue
        if line.startswith("- "):
            while i < len(lines) and lines[i].strip().startswith("- "):
                value = lines[i].strip()[2:].strip()
                i += 1
                while i < len(lines) and lines[i].strip() and not lines[i].strip().startswith(("- ", "|", "#", "> ")):
                    value += " " + lines[i].strip()
                    i += 1
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.06
                add_inline(p, value, size=9.35)
            continue
        if line.startswith("> "):
            value = line[2:].strip()
            i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.12)
            p.paragraph_format.right_indent = Inches(0.12)
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(5)
            add_inline(p, value, size=9.1)
            ppr = p._p.get_or_add_pPr()
            pborder = OxmlElement("w:pBdr")
            left = OxmlElement("w:left")
            left.set(qn("w:val"), "single")
            left.set(qn("w:sz"), "22")
            left.set(qn("w:space"), "6")
            left.set(qn("w:color"), TERRACOTTA)
            pborder.append(left)
            ppr.append(pborder)
            shd = OxmlElement("w:shd")
            shd.set(qn("w:fill"), CREAM)
            ppr.append(shd)
            continue
        text_parts = [line]
        i += 1
        while i < len(lines):
            next_line = lines[i].strip()
            if not next_line or next_line.startswith(("- ", "|", "> ")) or re.match(r"^#{1,6}\s", next_line):
                break
            text_parts.append(next_line)
            i += 1
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(5)
        p.paragraph_format.line_spacing = 1.07
        paragraph_text = " ".join(text_parts)
        if "Los pendientes P-02 a P-14 son:" in paragraph_text:
            p.paragraph_format.keep_with_next = True
        add_inline(p, paragraph_text, size=9.5)


def clear_paragraph_runs(paragraph):
    for child in list(paragraph._p):
        if child.tag != qn("w:pPr"):
            paragraph._p.remove(child)


def remove_title_border(doc, paragraph):
    ppr = paragraph._p.get_or_add_pPr()
    pborder = ppr.find(qn("w:pBdr"))
    if pborder is not None:
        ppr.remove(pborder)
    style = doc.styles["Title"]._element
    style_ppr = style.find(qn("w:pPr"))
    if style_ppr is not None:
        style_bdr = style_ppr.find(qn("w:pBdr"))
        if style_bdr is not None:
            style_ppr.remove(style_bdr)


def set_page_header(doc):
    for section in doc.sections:
        for table in section.header.tables:
            for row in table.rows:
                for cell in row.cells:
                    for p in cell.paragraphs:
                        if "ZA-ALC-001" in p.text:
                            clear_paragraph_runs(p)
                            r = p.add_run("ZHUD ALELÍ  |  ZA-ALC-GEN-001  |  ALCANCE GENERAL")
                            set_run(r, bold=True, size=8, color=GREEN)
                            p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def date_display(value):
    y, m, d = [int(part) for part in value.split("-")]
    months = ["enero", "febrero", "marzo", "abril", "mayo", "junio", "julio", "agosto", "septiembre", "octubre", "noviembre", "diciembre"]
    return f"{d} de {months[m - 1]} de {y}"


def main():
    source_text = SOURCE.read_text(encoding="utf-8-sig")
    metadata, source_body = parse_frontmatter(source_text)
    clean_body = re.sub(r"<style>.*?</style>", "", source_body, flags=re.S)
    clean_body = re.sub(r"^>\s*\*\*\[LOGO:.*?\n", "", clean_body, flags=re.M)
    doc = Document(TEMPLATE)
    doc.core_properties.title = "Alcance general del proyecto — Zhud Alelí v1.3"
    doc.core_properties.subject = "ZA-ALC-GEN-001 — Alcance general del proyecto"
    doc.core_properties.author = "Franklin Acero"
    doc.core_properties.keywords = "Zhud Alelí, alcance, requisitos, versión 1.3"
    doc.core_properties.comments = "Edición preparada para revisión del tesista; no aprobada en esta edición."
    section = doc.sections[0]
    section.page_width = Inches(8.27)
    section.page_height = Inches(11.69)
    section.left_margin = Inches(0.70)
    section.right_margin = Inches(0.70)
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.60)
    section.different_first_page_header_footer = True
    normal = doc.styles["Normal"]
    normal.font.name = "Heebo Light"
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Heebo Light")
    normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Heebo Light")
    normal.font.size = Pt(9.5)
    normal.font.color.rgb = RGBColor.from_string(INK)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.07
    for name, color, size in (("Heading 1", GREEN, 15), ("Heading 2", BURGUNDY, 12), ("Heading 3", BURGUNDY, 10.5)):
        style = doc.styles[name]
        style.font.name = "Georgia"
        style._element.rPr.rFonts.set(qn("w:ascii"), "Georgia")
        style._element.rPr.rFonts.set(qn("w:hAnsi"), "Georgia")
        style.font.color.rgb = RGBColor.from_string(color)
        style.font.size = Pt(size)
        style.font.bold = True
        style.paragraph_format.keep_with_next = True
    body = doc._element.body
    current = list(body)
    # Preserve cover layout/image/title/version and section properties; discard stale module content.
    preserved = current[:7]
    sect_pr = next((node for node in current if node.tag == qn("w:sectPr")), None)
    for node in current:
        if node not in preserved and node is not sect_pr:
            body.remove(node)
    title = doc.paragraphs[5]
    clear_paragraph_runs(title)
    title_run = title.add_run("Alcance general del proyecto")
    set_run(title_run, bold=True, size=25, color=GREEN)
    title_run.font.name = "Georgia"
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Georgia")
    title_run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Georgia")
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    remove_title_border(doc, title)
    version = doc.paragraphs[6]
    clear_paragraph_runs(version)
    version.alignment = WD_ALIGN_PARAGRAPH.CENTER
    vr = version.add_run(f"Versión {metadata.get('version', '1.3')}  |  {date_display(metadata['fecha'])}")
    set_run(vr, bold=True, size=11, color=BURGUNDY)
    # New content starts on a separate page after the branded cover.
    pagebreak = doc.add_paragraph()
    pagebreak.add_run().add_break(WD_BREAK.PAGE)
    set_page_header(doc)
    metadata_rows = [
        ["CONTROL DOCUMENTAL — ZHUD ALELÍ", ""],
        ["Código del documento", metadata.get("codigo", "ZA-ALC-GEN-001")],
        ["Módulo / alcance", metadata.get("modulo", "Alcance general del proyecto")],
        ["Fase / entregable", metadata.get("fase", "Requisitos y diseño — Sprint 1")],
        ["Versión / fecha", f"v{metadata.get('version', '1.3')} — {date_display(metadata['fecha'])}"],
        ["Autor / analista", metadata.get("autor", "Franklin Acero")],
        ["Estado", "Preparado para revisión del tesista; esta edición aún no está aprobada"],
    ]
    metadata_table = add_table(doc, metadata_rows, "metadata")
    metadata_table.cell(0, 0).merge(metadata_table.cell(0, 1))
    add_heading(doc, "Historial de revisiones", 2)
    all_lines = source_body.splitlines()
    hist = next(i for i, line in enumerate(all_lines) if line.strip() == "## Historial de revisiones")
    table_start = hist + 1
    while table_start < len(all_lines) and not all_lines[table_start].strip():
        table_start += 1
    revision_rows, _ = parse_table(all_lines, table_start)
    add_table(doc, revision_rows, "revisions")
    convert_body(doc, parse_source_body(clean_body))
    # Keep every paragraph editable and identify the version without marking it approved.
    doc.save(OUTPUT)
    print(f"created={OUTPUT}")
    print(f"bytes={OUTPUT.stat().st_size}")
    print(f"paragraphs={len(doc.paragraphs)} tables={len(doc.tables)}")


if __name__ == "__main__":
    main()
